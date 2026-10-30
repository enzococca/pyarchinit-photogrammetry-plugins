"""The catalogue, the drawings and the plugin: measures from geometry and model, room and space, exports, the sheets
(scale, what they say, a plan that is not a mirror image, DXF layers), the command line through the program's ``ppg
plugin``, the pane in the program's window."""
from __future__ import annotations

import csv
from pathlib import Path

import numpy as np
import pytest
from hulls import D, L, half_breadth, project

import pyarchinit_photogrammetry as ppg  # noqa: F401
from ppg_shipwreck import drawings, lines, ship, timbers


def _ready(tmp_path, heading=37.0, bow="end", water=False, **kw):
    doc, ch, ends = project(tmp_path, heading_deg=heading, **kw)
    ship.set_frame(ch, ship.Frame(ends[0], ends[1], bow))
    ship.set_stations(ch, ship.regular(2.0, 5, start=2.0))
    lines.build(ch, longitudinals=[0.0, 0.5], horizontals=[-0.5, -1.0], half_width=3.0)
    if water:
        lines.set_water(ch, -0.5, on=True)
    return doc, ch, ship.frame_of(ch)


def _world(frame, *points):
    return frame.to_world(np.array([[x, y, 0.0] for x, y in points]))[:, :2]


def test_measures_from_the_geometry_and_the_model(tmp_path):
    _doc, ch, frame = _ready(tmp_path)
    timbers.add(ch, "keel", "centre", "K", timbers.geometry("line", _world(frame, (2.0, 0.0), (10.0, 0.0))))
    s = half_breadth(6.0) * 0.5
    timbers.add(ch, "frame", "port", "F1", timbers.geometry("line", _world(frame, (6.0, 0.0), (6.0, s))), station="S3")
    timbers.add(ch, "plank", "port", "P", timbers.geometry("polygon", _world(frame, (3.0, 0.1), (5.0, 0.1), (5.0, 0.4), (3.0, 0.4))))
    timbers.add(ch, "fastening", "port", "N", timbers.geometry("point", _world(frame, (6.0, 0.5))))
    rows = {r["name"]: r for r in timbers.measure(ch)}
    assert rows["K"]["length"] > 8.0 and rows["K"]["z_min"] == pytest.approx(-D, abs=0.01)  # the keel follows the bottom
    # along the frame, on the curve of the section: z = D (y^2 / b^2 - 1) from y = 0 to b / 2
    b = half_breadth(6.0)
    y = np.linspace(0.0, s, 2000)
    truth = np.hypot(np.diff(y), np.diff(D * (y**2 / b**2 - 1.0))).sum()
    assert rows["F1"]["length"] == pytest.approx(truth, abs=0.01)
    assert rows["F1"]["width"] is None and rows["F1"]["x"] == pytest.approx(6.0)
    assert rows["P"]["length"] == pytest.approx(2.0, abs=1e-6) and rows["P"]["width"] == pytest.approx(0.3, abs=1e-6)
    assert rows["N"]["length"] is None
    assert rows["N"]["z_max"] == pytest.approx(D * (0.25 / b**2 - 1.0), abs=0.01)
    above = {r["name"]: r for r in timbers.measure(ch, datum=-D)}
    assert above["K"]["z_min"] == pytest.approx(0.0, abs=0.01)


def test_the_water_also_corrects_the_heights_of_the_timbers(tmp_path):
    _doc, ch, frame = _ready(tmp_path, water=True)
    timbers.add(ch, "keel", "centre", "K", timbers.geometry("line", _world(frame, (5.0, 0.0), (7.0, 0.0))))
    r = timbers.measure(ch)[0]
    assert r["z_min_water"] == pytest.approx(-0.5 - (-0.5 - r["z_min"]) * 1.34, abs=1e-6)


def test_room_and_space(tmp_path):
    _doc, ch, frame = _ready(tmp_path)
    for name, x, side in (("F1", 3.0, "port"), ("F2", 3.4, "port"), ("F3", 3.9, "port"), ("F1", 3.0, "starboard"), ("M", 3.2, "centre")):
        timbers.add(ch, "frame", side, name, timbers.geometry("point", _world(frame, (x, 0.3 if side == "port" else -0.3))))
    timbers.add(ch, "plank", "port", "P", timbers.geometry("point", _world(frame, (3.1, 0.3))))
    spacing = timbers.frame_spacing(timbers.measure(ch))
    assert spacing["port"]["spaces"] == pytest.approx([0.2, 0.2, 0.5])
    assert spacing["port"]["median"] == pytest.approx(0.2)
    assert [f["name"] for f in spacing["starboard"]["frames"]] == ["F1", "M"]
    stations = timbers.stations_from_catalogue(ch)
    assert {s["name"]: s["x"] for s in stations}["F1"] == pytest.approx(3.0)
    assert len(stations) == 4


def test_table_of_timbers_and_exports(tmp_path):
    import rasterio  # noqa: F401 - the program's stack
    from openpyxl import load_workbook

    from pyarchinit_photogrammetry.io import gpkg

    _doc, ch, frame = _ready(tmp_path)
    table = tmp_path / "t.csv"
    table.write_text("tipo;lato;nome;x1;y1;x2;y2;note\nframe;port;F36;4;0;4;0,5;rotta\nframe;starboard;F36;4;0;4;-0,5;\n", encoding="utf-8")
    assert timbers.import_csv(ch, table) == 2
    (tmp_path / "p.csv").write_text("kind,side,name,wkt\nkeelson,centre,KS,\"POINT (%f %f)\"\n" % tuple(_world(frame, (5.0, 0.0))[0]))
    timbers.import_csv(ch, tmp_path / "p.csv")
    rows = timbers.measure(ch)
    assert rows[0]["notes"] == "rotta" and rows[0]["x"] == pytest.approx(4.0)
    with open(timbers.export_any(ch, tmp_path / "a.csv")) as f:
        assert len(list(csv.DictReader(f))) == 3
    book = load_workbook(timbers.export_any(ch, tmp_path / "a.xlsx"))
    assert book.sheetnames == ["Legni", "Maglia"] and book["Legni"]["B2"].value == "ordinata"
    path = timbers.export_any(ch, tmp_path / "a.gpkg")
    assert len(gpkg.read_table(path, "timbers_line")) == 2 and len(gpkg.read_table(path, "timbers_point")) == 1
    with pytest.raises(ship.NotReady):
        timbers.export_any(ch, tmp_path / "a.txt")
    (tmp_path / "bad.csv").write_text("tipo;nome;x;y\nvela;A;1;2\n", encoding="utf-8")
    with pytest.raises(ship.NotReady, match="Riga 2"):
        timbers.import_csv(ch, tmp_path / "bad.csv")
    assert len(timbers.timbers_of(ch)) == 3  # a table with a bad line adds nothing


def test_the_drawings_are_written_with_their_notes(tmp_path):
    _doc, ch, frame = _ready(tmp_path, water=True)
    timbers.add(ch, "frame", "port", "F1", timbers.geometry("line", _world(frame, (6.0, 0.0), (6.0, 0.5))))
    ch.meta[ship.KEY]["title"] = {"site": "Prova", "vessel": "Scafo A"}
    written = drawings.write_all(ch)
    names = {p.name for p in written}
    for kind in drawings.KINDS:
        assert {f"{kind}.svg", f"{kind}.pdf", f"{kind}.png", f"{kind}.dxf"} <= names
    assert all(p.parent == ch.data_dir / "shipwreck" / "drawings" for p in written)
    svg = (written[0].parent / "body_plan.svg").read_text()
    assert "1,34" in svg and "Prova" in svg and "carenatura" in svg
    d = drawings.make(ch, "body_plan")
    assert d.scale == 20
    # raw line kept under the corrected one, the level drawn
    assert {"STATIONS", "RAW", "LEVEL", "KEEL"} == {layer for layer, _p, _c in d.lines}
    import ezdxf

    doc = ezdxf.readfile(written[0].parent / "body_plan.dxf")
    used = {e.dxf.layer for e in doc.modelspace()}
    assert {"STATIONS", "RAW", "LEVEL", "TEXT"} <= used
    plan = ezdxf.readfile(written[0].parent / "plan.dxf")
    assert "TIMBERS" in {e.dxf.layer for e in plan.modelspace()}
    assert ch.meta[ship.KEY]["drawings"]["profile"]["water"] is True


def test_the_scale_is_exact_on_the_paper(tmp_path):
    _doc, ch, frame = _ready(tmp_path)
    d = drawings.make(ch, "body_plan", scale=50)
    fig = drawings.figure(d)
    ax = fig.axes[0]
    width_mm = ax.get_position().width * fig.get_size_inches()[0] * 25.4
    x0, x1, _y0, _y1 = d.extent
    assert width_mm == pytest.approx((x1 - x0) * 1000.0 / 50, rel=1e-6)


def test_a_plan_is_never_a_mirror_image(tmp_path):
    """On the sheet the left of x is up, whichever way port lies: the world's north of an x to the east."""
    for bow in ("start", "end"):
        (tmp_path / bow).mkdir()
        _doc, ch, frame = _ready(tmp_path / bow, heading=90.0, bow=bow)  # x to the east: north is up on the sheet
        d = drawings.make(ch, "half_breadth")
        assert d.north == pytest.approx((0.0, 1.0), abs=1e-9)
        world = np.array([[*frame.to_world(np.array([[6.0, 0.0, 0.0]]))[0][:2], 0.0]]) + [0.0, 1.0, 0.0]
        assert frame.to_ship(world)[0, 1] * frame.sign > 0


def test_a_chunk_with_a_stale_or_missing_input_is_told(tmp_path):
    _doc, ch, frame = _ready(tmp_path)
    ship.set_frame(ch, frame.flipped())
    with pytest.raises(ship.NotReady, match="ricalcolale"):
        drawings.write_all(ch)
    drawings_folder = ch.data_dir / "shipwreck" / "drawings"
    assert not drawings_folder.exists() or not list(drawings_folder.iterdir())


def test_a_drawing_with_no_lines_is_left_out(tmp_path):
    doc, ch, ends = project(tmp_path)
    ship.set_frame(ch, ship.Frame(ends[0], ends[1]))
    ship.set_stations(ch, [{"name": "S", "x": 6.0}])
    lines.build(ch, half_width=3.0)
    names = {p.stem for p in drawings.write_all(ch)}
    assert names == {"body_plan", "plan"}


def test_the_plan_over_the_orthophoto(tmp_path):
    import rasterio
    from rasterio.transform import from_origin

    _doc, ch, frame = _ready(tmp_path, heading=30.0)
    # an orthophoto whose east half is red: on the sheet the port (+y) and starboard sides of an x to the
    # north-north-east show which side is which
    centre = frame.to_world(np.array([[6.0, 0.0, 0.0]]))[0]
    cell, n = 0.02, 1000
    x0, y1 = centre[0] - n * cell / 2, centre[1] + n * cell / 2
    image = np.zeros((3, n, n), np.uint8)
    image[0, :, n // 2:] = 255
    image[1, :, : n // 2] = 255
    path = ch.data_dir / "orthomosaic.tif"
    with rasterio.open(path, "w", driver="GTiff", width=n, height=n, count=3, dtype="uint8", crs=ch.crs,
                       transform=from_origin(x0, y1, cell, cell)) as dst:
        dst.write(image)
    ch._set_product("orthomosaic", path)
    d = drawings.make(ch, "plan", scale=20)
    assert d.raster is not None
    rgba, (a, b, c, e) = d.raster
    rows, cols = rgba.shape[:2]
    # the point of the sheet at (6 m, +1 m on the sheet = left of x): west of the keel when x runs north-north-east
    col = int((6.0 - a) / (b - a) * cols)
    row = int((e - 1.0) / (e - c) * rows)
    assert rgba[row, col, 1] == 255 and rgba[row, col, 0] == 0
    row = int((e + 1.0) / (e - c) * rows)
    assert rgba[row, col, 0] == 255 and rgba[row, col, 1] == 0


def _plugin_folder(root: Path) -> Path:
    folder = root / "plugins"
    dist = folder / "shipwreck" / "ppg_shipwreck-0.1.0.dist-info"
    dist.mkdir(parents=True)
    (dist / "METADATA").write_text("Metadata-Version: 2.1\nName: ppg-shipwreck\nVersion: 0.1.0\n")
    (dist / "entry_points.txt").write_text("[pyarchinit_photogrammetry.plugins]\nshipwreck = ppg_shipwreck:ShipwreckPlugin\n")
    return folder


def test_command_line(tmp_path, monkeypatch):
    from typer.testing import CliRunner

    from pyarchinit_photogrammetry.cli import app

    doc, ch, ends = project(tmp_path, heading_deg=37.0)
    monkeypatch.setenv("PPG_PLUGINS_DIR", str(_plugin_folder(tmp_path)))
    monkeypatch.setenv("PPG_PLUGINS_ENVIRONMENT", "0")
    p = str(doc.path)

    def run(*args):
        result = CliRunner().invoke(app, ["plugin", "shipwreck", args[0], p, *args[1:]])
        assert result.exit_code == 0, result.output
        return result.output

    run("frame", "--start", ",".join(map(str, ends[0])), "--end", ",".join(map(str, ends[1])), "--bow", "end")
    run("stations", "--spacing", "2", "--count", "5", "--first", "2")
    run("sections", "--longitudinals", "0;0.5", "--horizontals", "-0.5", "--half-width", "3")
    run("water", "--level", "-0.5", "--on")
    table = tmp_path / "t.csv"
    table.write_text("kind,side,name,x,y\nframe,port,F36,6,0.5\n")
    assert "F36" in run("timbers", "--table", str(table))
    run("drawings", "--scale", "50")
    out = tmp_path / "legni.xlsx"
    run("export", "--out", str(out))
    assert out.exists()
    assert "sezioni" in run("status")
    again = ppg.Document.open(p).chunk
    assert len(ship.stations_of(again)) == 5 and lines.is_on(again) and len(timbers.timbers_of(again)) == 1
    assert (again.data_dir / "shipwreck" / "drawings" / "body_plan.pdf").exists()
    bad = CliRunner().invoke(app, ["plugin", "shipwreck", "water", p])
    assert bad.exit_code == 1


def test_the_pane_in_the_window(tmp_path):
    from PySide6.QtWidgets import QApplication

    from ppg_shipwreck import ShipwreckPlugin
    from pyarchinit_photogrammetry.app import app
    from pyarchinit_photogrammetry.gui import dialogs
    from pyarchinit_photogrammetry.gui import main_window as mwmod
    from pyarchinit_photogrammetry.plugins import Context

    QApplication.instance() or QApplication([])
    doc, ch, ends = project(tmp_path, heading_deg=37.0)
    app.document = doc
    w = mwmod.MainWindow(doc, autoconnect=False, check_updates=False)
    w.runner.show_dialog = False
    plugin = ShipwreckPlugin()
    answers = {}
    original = dialogs.ParamDialog.ask
    dialogs.ParamDialog.ask = classmethod(lambda cls, parent, title, fields, intro="": answers.pop("next", None))
    try:
        context = Context("shipwreck", tmp_path / "plugin", window=w)
        plugin.activate(context)
        pane = plugin.pane
        assert w.panes.groups["right"].dock(pane) is not None
        assert "Non impostato" in pane.states["frame"].text()
        # a step that misses what it needs says so
        messages = []
        import ppg_shipwreck

        plugin._message = lambda kind, text: messages.append(text)
        pane.requested.emit("sections")
        assert "sistema della nave" in messages[-1]
        # the keel traced on the view with the program's tool
        pane.requested.emit("frame_pick")
        assert w.ortho_view.tr is not None  # an empty frame around the model: there is no raster
        middle = ch.to_world(np.array([[1.0, 0.0, 0.0], [5.0, 0.0, 0.0]]))
        plugin._set_picked(ch, middle[:, :2])
        frame = ship.frame_of(ch)
        assert frame.length == pytest.approx(8.0)
        assert "Chiglia 8,00 m" in pane.states["frame"].text()
        answers["next"] = {"spacing": 2.0, "count": 3, "start": 0.0}
        pane.requested.emit("stations_regular")
        assert "3 stazioni" in pane.states["stations"].text()
        answers["next"] = {"longitudinals": "0", "horizontals": "-0,5", "half_width": 3.0}
        pane.requested.emit("sections")
        w.runner.wait()
        assert "5 sezioni" in pane.states["sections"].text()
        answers["next"] = {"level": -0.5, "n": 1.34, "on": True}
        pane.requested.emit("water")
        assert "correzione accesa" in pane.states["water"].text()
        answers["next"] = {"kind": "frame", "side": "port", "name": "F36", "station": "", "notes": ""}
        plugin._timber("line", ch.to_world(np.array([[2.0, 0.0, 0.0], [2.0, 0.2, 0.0]]))[:, :2])
        assert pane.table.rowCount() == 1 and pane.table.item(0, 3).text() == "F36"
        assert ppg_shipwreck.LAYER in w.ortho_view.layers
        labels = [i.get("label") for i in plugin._tree_items(ch)[0]["children"]]
        assert any("Legni (1)" in x for x in labels) and any("Sezioni" in x for x in labels)
        # deleted on the view, gone from the catalogue
        plugin._features_deleted([{"id": "T1"}])
        assert not timbers.timbers_of(ch) and pane.table.rowCount() == 0
    finally:
        dialogs.ParamDialog.ask = original
        context.release()
        w.setWindowModified(False)
        w.close()
