"""The catalogue of the timbers: records in ``chunk.meta["shipwreck"]["timbers"]`` with their geometry in the
coordinates of the chunk (a line, a polygon or a point drawn on the raster, or a CSV), and the measures computed from
the geometry and the model: the length along the timber, the width across it, the heights above a datum. The spacing
of the frames (room and space) comes from their x. Exports: CSV, XLSX, GeoPackage."""
from __future__ import annotations

import csv
from pathlib import Path

import numpy as np

from . import lines, ship
from .texts import num, t

KINDS = ("keel", "keelson", "frame", "plank", "ceiling", "stringer", "fastening", "other")
SIDES = ("port", "starboard", "centre")
KIND_TITLES = {"keel": ("chiglia", "keel"), "keelson": ("paramezzale", "keelson"), "frame": ("ordinata", "frame"),
               "plank": ("corso di fasciame", "plank"), "ceiling": ("fasciame interno", "ceiling"),
               "stringer": ("serretta", "stringer"), "fastening": ("elemento di fissaggio", "fastening"),
               "other": ("altro", "other")}
SIDE_TITLES = {"port": ("babordo", "port"), "starboard": ("tribordo", "starboard"), "centre": ("centro", "centre")}
COLUMNS = ("id", "kind", "side", "name", "station", "x", "y", "length", "width", "z_min", "z_max", "z_min_water",
           "z_max_water", "notes", "warning")
SAMPLE = 0.02  # metres between the points read along a timber
MAX_SAMPLES = 2000
SIDE_TOLERANCE = 0.05  # metres from the keel line within which a timber lies on neither side


def kind_title(kind: str) -> str:
    return t(*KIND_TITLES.get(kind, (kind, kind)))


def side_title(side: str) -> str:
    return t(*SIDE_TITLES.get(side, (side, side)))


def timbers_of(chunk) -> list[dict]:
    return list(chunk.meta.get(ship.KEY, {}).get("timbers", []))


def geometry(kind: str, xy) -> dict:
    """A geometry of a record: ``kind`` "line", "polygon" or "point", ``xy`` in the coordinates of the chunk."""
    xy = np.asarray(xy, float).reshape(-1, 2)
    needed = {"line": 2, "polygon": 3, "point": 1}[kind]
    if len(xy) < needed:
        raise ship.NotReady(t("Geometria con troppi pochi vertici.", "A geometry with too few vertices."))
    return {"type": kind, "xy": xy.tolist()}


def _code(value: str, titles: dict, what: tuple) -> str:
    """The code of a kind or a side, also written in Italian."""
    text = value.strip().lower()
    for code, names in titles.items():
        if text == code or text in names:
            return code
    raise ship.NotReady(t(f"{what[0]} sconosciuto: {value}. Valori: {', '.join(titles)}.",
                          f"Unknown {what[1]}: {value}. Values: {', '.join(titles)}."))


def add(chunk, kind: str, side: str, name: str, geom: dict, station: str = "", notes: str = "") -> dict:
    kind, side = _code(kind, KIND_TITLES, ("Tipo", "kind")), _code(side, SIDE_TITLES, ("Lato", "side"))
    store = ship.state(chunk).setdefault("timbers", [])
    n = 1 + max((int(r["id"][1:]) for r in store if r["id"][1:].isdigit()), default=0)
    record = {"id": f"T{n}", "kind": kind, "side": side, "name": str(name), "station": str(station),
              "geometry": geom, "notes": str(notes)}
    store.append(record)
    return record


def remove(chunk, ids) -> int:
    store = ship.state(chunk).setdefault("timbers", [])
    before = len(store)
    store[:] = [r for r in store if r["id"] not in set(ids)]
    return before - len(store)


def _value(row, header, names):
    i = next((header.index(n) for n in names if n in header), None)
    return row[i].strip() if i is not None and i < len(row) else ""


def import_csv(chunk, path) -> int:
    """Records from a table with a header: ``kind`` (or ``tipo``), ``side`` (``lato``), ``name`` (``nome``),
    ``station`` (``stazione``), ``notes`` (``note``) and the geometry in ship coordinates: ``x, y`` (a point) or
    ``x1, y1, x2, y2`` (a line, a timber measured in the field); or ``wkt`` in the coordinates of the chunk."""
    frame = ship.need_frame(chunk)
    header, rows = ship.rows_of(path)
    if not header:
        raise ship.NotReady(t("La tabella dei legni ha bisogno dell'intestazione.", "The table of timbers needs a header."))
    parsed = []
    for i, row in rows:
        try:
            parsed.append(_record(row, header, frame))
        except (ValueError, IndexError, ship.NotReady, KeyError) as exc:
            raise ship.NotReady(t(f"Riga {i} della tabella dei legni: {exc}", f"Line {i} of the table of timbers: {exc}")) from None
    for args in parsed:
        add(chunk, *args)
    return len(parsed)


def _record(row, header, frame):
    wkt = _value(row, header, ("wkt",))
    if wkt:
        import shapely

        shape = shapely.from_wkt(wkt)
        names = {"LineString": "line", "Polygon": "polygon", "Point": "point"}
        xy = (np.asarray(shape.exterior.coords)[:-1] if shape.geom_type == "Polygon" else np.asarray(shape.coords))[:, :2]
        geom = geometry(names[shape.geom_type], xy)
    else:
        if _value(row, header, ("x1",)):
            ends = [[ship.number(_value(row, header, (k,))) for k in keys] for keys in (("x1", "y1"), ("x2", "y2"))]
            kind = "line"
        else:
            ends = [[ship.number(_value(row, header, ("x",))), ship.number(_value(row, header, ("y",)))]]
            kind = "point"
        world = frame.to_world(np.column_stack([np.asarray(ends), np.zeros(len(ends))]))
        geom = geometry(kind, world[:, :2])
    kind = _code(_value(row, header, ("kind", "tipo")) or "other", KIND_TITLES, ("Tipo", "kind"))
    side = _code(_value(row, header, ("side", "lato")) or "centre", SIDE_TITLES, ("Lato", "side"))
    return kind, side, _value(row, header, ("name", "nome")), geom, _value(row, header, ("station", "stazione")), \
        _value(row, header, ("notes", "note"))


def _samples(geom: dict, frame: ship.Frame) -> tuple[np.ndarray, str]:
    """The ship (x, y) of the points read on the model: along a line, over a polygon, or the point."""
    xy = np.asarray(geom["xy"], float)
    s = frame.to_ship(np.column_stack([xy, np.zeros(len(xy))]))[:, :2]
    if geom["type"] == "point":
        return s, "point"
    if geom["type"] == "line":
        seg = np.hypot(*np.diff(s, axis=0).T)
        along = np.r_[0.0, np.cumsum(seg)]
        count = int(min(MAX_SAMPLES, max(2, along[-1] / SAMPLE)))
        d = np.linspace(0.0, along[-1], count)
        return np.column_stack([np.interp(d, along, s[:, 0]), np.interp(d, along, s[:, 1])]), "line"
    import shapely

    polygon = shapely.Polygon(s)
    x0, y0, x1, y1 = polygon.bounds
    step = max(SAMPLE, np.sqrt((x1 - x0) * (y1 - y0) / MAX_SAMPLES))
    gx, gy = np.meshgrid(np.arange(x0, x1 + step, step), np.arange(y0, y1 + step, step))
    inside = shapely.contains_xy(polygon, gx.ravel(), gy.ravel())
    return np.vstack([s, np.column_stack([gx.ravel()[inside], gy.ravel()[inside]])]), "polygon"


def _polygon_sides(geom: dict, frame: ship.Frame) -> tuple[float, float]:
    """Long and short side of the smallest rectangle around a polygon."""
    import shapely

    xy = np.asarray(geom["xy"], float)
    box = shapely.minimum_rotated_rectangle(shapely.Polygon(xy - xy.mean(axis=0)))
    c = np.asarray(box.exterior.coords)
    a, b = np.hypot(*(c[1] - c[0])), np.hypot(*(c[2] - c[1]))
    return float(max(a, b)), float(min(a, b))


def _side_warning(side: str, y: float) -> str:
    """Words when a record said port or starboard lies on the other side of the frame."""
    wrong = (side == "port" and y < -SIDE_TOLERANCE) or (side == "starboard" and y > SIDE_TOLERANCE)
    if not wrong:
        return ""
    return t(f"segnato {side_title(side)} ma sta dall'altro lato (y {num(y, 2)})",
             f"said {side_title(side)} but lies on the other side (y {num(y, 2)})")


def _length_over(s: np.ndarray, z: np.ndarray, ok: np.ndarray, reach: float) -> float:
    """The length of a line over the model, with the heights read ``reach`` apart (a few edges of the mesh):
    read closer, the noise of the surface adds up as length. Plan length where the model is not under it."""
    filled = np.interp(np.arange(len(z)), np.flatnonzero(ok), z[ok]) if ok.any() else np.zeros(len(z))
    plan = np.r_[0.0, np.cumsum(np.hypot(*np.diff(s, axis=0).T))]
    count = max(2, int(round(plan[-1] / max(reach, SAMPLE))) + 1)
    at = np.linspace(0.0, plan[-1], count)
    return float(np.hypot(np.diff(at), np.diff(np.interp(at, plan, filled))).sum())


def measure(chunk, datum: float | None = None, progress=None) -> list[dict]:
    """A row for each record, ``COLUMNS``: ``x``, ``y`` the mean position in ship coordinates; ``length`` along the
    timber over the model (for a polygon: the long side of the rectangle around it, as ``width`` the short one; a
    line has no width); ``z_min``, ``z_max`` the heights of the model on it above ``datum`` (ship z; default 0); with
    the water on, ``z_min_water``, ``z_max_water`` with the correction. Empty where the model is not under it, and
    all the heights without a model (lengths in plan). ``warning``: the record's side against the sign of its y."""
    frame = ship.need_frame(chunk)
    records = timbers_of(chunk)
    if not records:
        return []
    model = chunk.product_path("model")
    surface = ship.Surface(lines.ship_mesh(chunk, frame).vertices) if model is not None and model.exists() else None
    water = lines.water_of(chunk) if lines.is_on(chunk) else None
    datum = float(datum or 0.0)
    rows = []
    for r in records:
        s, kind = _samples(r["geometry"], frame)
        z = surface.at(s) if surface is not None else np.full(len(s), np.nan)
        ok = np.isfinite(z)
        row = {k: r.get(k, "") for k in ("id", "kind", "side", "name", "station", "notes")}
        row["warning"] = _side_warning(r["side"], float(s[:, 1].mean()))
        row.update(x=float(s[:, 0].mean()), y=float(s[:, 1].mean()), length=None, width=None, z_min=None, z_max=None,
                   z_min_water=None, z_max_water=None)
        if kind == "polygon":
            row["length"], row["width"] = _polygon_sides(r["geometry"], frame)
        elif kind == "line":
            row["length"] = _length_over(s, z, ok, surface.reach if surface is not None else SAMPLE)
        if ok.any():
            row["z_min"], row["z_max"] = float(z[ok].min()) - datum, float(z[ok].max()) - datum
            if water:
                true = lines.corrected(z[ok], water["level"], water["n"])
                row["z_min_water"], row["z_max_water"] = float(true.min()) - datum, float(true.max()) - datum
        rows.append(row)
    return rows


def frame_spacing(rows: list[dict]) -> dict:
    """Room and space from the frames' x: for each side (a frame on the centre counts for both), the frames in order
    with the distance from the one before, and the median. ``{side: {"frames": [...], "spaces": [...], "median"}}``."""
    out = {}
    for side in ("port", "starboard"):
        frames = sorted((r for r in rows if r["kind"] == "frame" and r["side"] in (side, "centre")), key=lambda r: r["x"])
        if not frames:
            continue
        spaces = np.diff([f["x"] for f in frames]).tolist()
        out[side] = {"frames": [{"id": f["id"], "name": f["name"], "x": f["x"]} for f in frames], "spaces": spaces,
                     "median": float(np.median(spaces)) if spaces else None}
    return out


def stations_from_catalogue(chunk) -> list[dict]:
    """A station at the x of each frame of the catalogue, named after it; frames of the same name (the two
    sides of one frame) at their mean x."""
    rows = [r for r in measure(chunk) if r["kind"] == "frame"]
    if not rows:
        raise ship.NotReady(t("Il catalogo non ha ordinate.", "The catalogue has no frames."))
    names: dict[str, list[float]] = {}
    for r in rows:
        names.setdefault(r["name"] or r["id"], []).append(r["x"])
    return [{"name": n, "x": float(np.mean(x))} for n, x in names.items()]


# ------------------------------------------------------------------------------------- exports
def _cell(value):
    return "" if value is None else (round(value, 4) if isinstance(value, float) else value)


def _titled(row: dict) -> dict:
    return {**row, "kind": kind_title(row["kind"]), "side": side_title(row["side"])}


def export_csv(chunk, path, datum: float | None = None) -> Path:
    rows = measure(chunk, datum)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        for row in rows:
            w.writerow([_cell(row[c]) for c in COLUMNS])
    return Path(path)


def export_xlsx(chunk, path, datum: float | None = None) -> Path:
    """Two sheets: the timbers, the spacing of the frames."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    rows = measure(chunk, datum)
    book = Workbook()
    sheet = book.active
    sheet.title = t("Legni", "Timbers")
    sheet.append(list(COLUMNS))
    for row in rows:
        row = _titled(row)
        sheet.append([_cell(row[c]) for c in COLUMNS])
    second = book.create_sheet(t("Maglia", "Room and space"))
    second.append([t("lato", "side"), t("ordinata", "frame"), "x", t("distanza dalla precedente (m)", "from the one before (m)")])
    for side, s in frame_spacing(rows).items():
        for i, f in enumerate(s["frames"]):
            second.append([side_title(side), f["name"] or f["id"], round(f["x"], 4), round(s["spaces"][i - 1], 4) if i else None])
    for ws in (sheet, second):
        for cell in ws[1]:
            cell.font = Font(bold=True)
    book.save(path)
    return Path(path)


def _has_table(path, table: str) -> bool:
    from pyarchinit_photogrammetry.io import gpkg

    if not gpkg.is_geopackage(path):
        return False
    con = gpkg.connect(path)
    try:
        return con.execute("SELECT 1 FROM gpkg_contents WHERE table_name = ?", (table,)).fetchone() is not None
    finally:
        con.close()


def export_gpkg(chunk, path, datum: float | None = None) -> Path:
    """A table per kind of geometry, in the coordinates of the chunk, with the measures as columns."""
    import shapely

    from pyarchinit_photogrammetry.io import gpkg

    rows = {r["id"]: r for r in measure(chunk, datum)}
    columns = {c: ("REAL" if c in ("x", "y", "length", "width", "z_min", "z_max", "z_min_water", "z_max_water") else "TEXT")
               for c in COLUMNS}
    makers = {"line": ("LINESTRING", shapely.LineString), "polygon": ("POLYGON", shapely.Polygon),
              "point": ("POINT", lambda xy: shapely.Point(xy[0]))}
    records = timbers_of(chunk)
    if not records:
        raise ship.NotReady(t("Il catalogo è vuoto.", "The catalogue is empty."))
    for kind, (geometry_type, make) in makers.items():
        chosen = [(make(r["geometry"]["xy"]), rows[r["id"]]) for r in records if r["geometry"]["type"] == kind]
        if chosen:
            gpkg.write_table(path, f"timbers_{kind}", geometry_type, columns, chosen, chunk.crs)
        elif _has_table(path, f"timbers_{kind}"):  # an earlier export: a kind no longer in the catalogue is emptied
            gpkg.write_table(path, f"timbers_{kind}", geometry_type, columns, [], chunk.crs)
    return Path(path)


def export_any(chunk, path, datum: float | None = None) -> Path:
    """By the suffix of ``path``: .csv, .xlsx, .gpkg."""
    writers = {".csv": export_csv, ".xlsx": export_xlsx, ".gpkg": export_gpkg}
    writer = writers.get(Path(path).suffix.lower())
    if writer is None:
        raise ship.NotReady(t("Formati: .csv, .xlsx, .gpkg.", "Formats: .csv, .xlsx, .gpkg."))
    return writer(chunk, path, datum)
